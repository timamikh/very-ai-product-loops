#!/usr/bin/env python3
"""to-document · one-pager profile → a formatted .docx product brief for decksmith.

A base-adapter render: neutral, house-agnostic styling. Content is a projection of the instance
(passport / analysis / strategy / strategic-plan / sprint-plan + registers); the artifacts stay the
source of truth. Re-run to regenerate:  python3 render_brief.py
Requires: python-docx  (pip install python-docx)
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.normpath(os.path.join(BASE, "..", "deliverables"))
INK   = RGBColor(0x14, 0x17, 0x1D)
ACCENT= RGBColor(0x1D, 0x46, 0xD6)
MUTED = RGBColor(0x5B, 0x64, 0x70)
WARN  = RGBColor(0xC0, 0x2A, 0x22)
FONT  = "Calibri"

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.9); s.bottom_margin = Cm(1.9)
    s.left_margin = Cm(2.1); s.right_margin = Cm(2.1)
normal = doc.styles["Normal"]
normal.font.name = FONT; normal.font.size = Pt(10.5); normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.18

def rule(color=ACCENT, size=14, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); btm = OxmlElement("w:bottom")
    btm.set(qn("w:val"), "single"); btm.set(qn("w:sz"), str(size)); btm.set(qn("w:space"), "1")
    btm.set(qn("w:color"), "%02X%02X%02X" % (color[0], color[1], color[2]))
    bdr.append(btm); pPr.append(bdr)

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(10)
    r.font.color.rgb = ACCENT; r.font.name = FONT
    rPr = r._element.get_or_add_rPr(); sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "30"); rPr.append(sp)

def body(text):
    doc.add_paragraph().add_run(text)

def bullet(label, text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
    if label:
        r = p.add_run(label + " — "); r.bold = True; r.font.color.rgb = INK
    p.add_run(text)

def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexfill); tcPr.append(sh)

def setcell(cell, text, bold=False, color=INK, size=9.5, white=False):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if white else color

# ---- header ----
t = doc.add_paragraph(); t.paragraph_format.space_after = Pt(0)
rt = t.add_run("Decksmith"); rt.bold = True; rt.font.size = Pt(30); rt.font.color.rgb = INK
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(1)
rs = sub.add_run("Concept Brief — AI that generates editable, well-designed native decks")
rs.font.size = Pt(12); rs.font.color.rgb = MUTED
meta = doc.add_paragraph()
rm = meta.add_run("Status: concept-viability   ·   July 2026   ·   fictional sample")
rm.font.size = Pt(8.5); rm.font.color.rgb = MUTED; rm.font.name = "Consolas"
rule()

heading("The concept")
body("An AI tool that generates a full presentation and exports it as a native .pptx / .key file "
     "where every object stays genuinely editable and genuinely designed — not a picture of a slide, "
     "and not an obvious template.")

heading("The problem — the job")
body("A salesperson prepping a client pitch under time pressure wants a deck that's ready to send, "
     "not a draft to rebuild. Today's AI decks fail three ways:")
bullet("Looks templated", "generic, off-brand output that reads “made by AI” at a glance (P1).")
bullet("Wrong structure", "passable visuals wrapped around the wrong argument and framing (P2).")
bullet("Can't edit it", "export lands as flattened images, so the deck gets redone by hand — the killer.")

heading("The opportunity — the white space")
p = doc.add_paragraph()
p.add_run("Every AI deck tool is web-first. ")
rr = p.add_run("30–40% of slides are flattened to images when Gamma exports to .pptx")
rr.bold = True; rr.font.color.rgb = WARN
p.add_run(" [sourced: market-research]. The unclaimed ground is native fidelity — a downloaded file "
          "that stays fully editable in PowerPoint and Keynote.")

heading("Market")
bullet("Size", "~$750M obtainable SAM, bottom-up from reachable segments [assumption].")
bullet("Who's here", "Gamma, Microsoft Copilot, Canva, Beautiful.ai, Pitch, Tome.")
bullet("The threat", "Copilot and Canva are entering the same wedge — so the moat must be the "
                     "fidelity + design engine, not the app around it.")

heading("How we win")
bullet("Beachhead", "salespeople prepping a client pitch under deadline — a sharp, reachable job.")
bullet("Positioning", "“actually-editable + designed native decks” — a wedge strong enough to "
                      "beat the PowerPoint habit.")
bullet("Pricing", "$15–25 / mo standalone — an open willingness-to-pay hypothesis, not a set price.")

heading("How we'll know")
bullet("The gate metric", "edit-fidelity ≥ 90% of exported objects natively editable (M-edit-fidelity) "
                          "— the concept-proving read.")
bullet("Guardrail", "COGS ≤ $1 / deck (LLM inference is a cost line). Red line: no flatten-to-image "
                    "shortcut — it would fake the entire wedge.")

heading("The bet & the test")
body("A 6-week Concept Test with one learning goal: can we build native edit-fidelity at scale (H-001)? "
     "Sprint 1 ships a thin export slice, a fidelity harness, and a cheap demand probe. Decision rule is "
     "pre-registered:")

tbl = doc.add_table(rows=1, cols=4); tbl.alignment = WD_TABLE_ALIGNMENT.LEFT; tbl.autofit = True
for c, txt in zip(tbl.rows[0].cells, ["Hypothesis", "What it claims", "Test / threshold", "Status"]):
    setcell(c, txt, bold=True, white=True); shade(c, "1D46D6")
rows = [
    ("H-001", "The engine hits native edit-fidelity at scale", "20 test decks; ≥90% build · <70% pivot", "testing"),
    ("H-005", "A native editable+designed deck is a real unmet need", "Demand probe B-01/B-03; ≥10 qualified trials", "testing"),
]
for i, (a, b, c, d) in enumerate(rows):
    cells = tbl.add_row().cells
    setcell(cells[0], a, bold=True, color=ACCENT); setcell(cells[1], b)
    setcell(cells[2], c); setcell(cells[3], d, bold=True)
    if i % 2 == 0:
        for cc in cells: shade(cc, "F3F5F9")

rule(space_after=4)
foot = doc.sections[0].footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
rf = foot.add_run("Fictional sample · rendered from very-ai-product-loops (to-document) · "
                  "source of truth: examples/decksmith/ artifacts + registers")
rf.font.size = Pt(7.5); rf.font.color.rgb = MUTED; rf.font.name = "Consolas"

out = os.path.join(OUT_DIR, "concept-brief.docx")
doc.save(out)
print("saved", out)
