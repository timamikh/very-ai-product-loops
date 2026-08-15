#!/usr/bin/env python3
"""to-table — GENERIC renderer: markdown tables from any instance -> one styled .xlsx (a tab per
source) or CSVs. Instance-agnostic: no product data lives here, only the rendering logic — so it
travels with the framework. Requires openpyxl (a table library is why this is code, not instructions).

Usage:
  render.py INSTANCE_DIR [--section "relpath[#anchor]:Tab Name"]... [--no-registers]
                         [--out FILE] [--format xlsx|csv] [--outdir DIR]

By default renders the three registers (hypotheses, risks, metric-tree) as three tabs; add artifact
tables with --section. Each source contributes the first markdown table (after #anchor if given).

Example (decksmith):
  render.py examples/decksmith \\
    --section "4-strategic-plan.md#global-hypotheses:Strategic plan" \\
    --section "6-sprint-plan.md#backlog:Sprint plan" \\
    --out examples/decksmith/deliverables/decksmith-registers-and-plan.xlsx
"""
import argparse, csv, os, re, sys

DEFAULT_REGISTERS = [
    ("registers/hypotheses.md", None, "Hypotheses"),
    ("registers/risks.md", None, "Risks"),
    ("registers/metric-tree.md", None, "Metric tree"),
]

def clean(s):
    s = s.replace("**", "").replace("`", "")
    return re.sub(r"\s+", " ", s).strip()

def parse_md_table(text, anchor=None):
    """Return (header, rows) for the first pipe-table (optionally after a {#anchor} heading)."""
    lines = text.splitlines()
    i = 0
    if anchor:
        for j, l in enumerate(lines):
            if ("{#" + anchor + "}") in l:
                i = j + 1
                break
        else:
            raise ValueError("anchor #%s not found" % anchor)
    while i < len(lines):
        l = lines[i].strip()
        nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
        if l.startswith("|") and nxt and set(nxt) <= set("|:- ") and "-" in nxt:
            header = [clean(c) for c in l.strip("|").split("|")]
            rows, i = [], i + 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([clean(c) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            return header, rows
        i += 1
    raise ValueError("no markdown table found")

def load_sources(instance, sections, use_registers):
    specs = list(DEFAULT_REGISTERS) if use_registers else []
    for s in sections:
        left, _, tab = s.rpartition(":")
        path, _, anchor = left.partition("#")
        specs.append((path, anchor or None, tab))
    out = []
    for path, anchor, tab in specs:
        full = os.path.join(instance, path)
        header, rows = parse_md_table(open(full, encoding="utf-8").read(), anchor)
        src = path + ("#" + anchor if anchor else "")
        out.append((tab, header, rows, src))
    return out

def write_xlsx(sheets, out, stamp):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    ACCENT, ZEBRA, INK, MUTED = "1D46D6", "F3F5F9", "14171D", "5B6470"
    thin = Side(style="thin", color="E4E1DA"); border = Border(thin, thin, thin, thin)
    wb = Workbook(); wb.remove(wb.active)
    for tab, header, rows, src in sheets:
        ws = wb.create_sheet(tab[:31]); n = len(header)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        cap = ws.cell(1, 1, "source: %s · rendered %s (to-table)" % (src, stamp))
        cap.font = Font(name="Calibri", italic=True, color=MUTED, size=8.5)
        ws.row_dimensions[1].height = 16
        for j, h in enumerate(header, 1):
            c = ws.cell(2, j, h); c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=ACCENT)
            c.alignment = Alignment(vertical="center", wrap_text=True); c.border = border
        ws.row_dimensions[2].height = 20
        for i, row in enumerate(rows):
            for j in range(n):
                val = row[j] if j < len(row) else ""
                c = ws.cell(i + 3, j + 1, val); c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(name="Calibri", bold=(j == 0), color=ACCENT if j == 0 else INK, size=10)
                if i % 2 == 1:
                    c.fill = PatternFill("solid", fgColor=ZEBRA)
        for j in range(n):
            longest = max([len(header[j])] + [len(r[j]) for r in rows if j < len(r)])
            ws.column_dimensions[get_column_letter(j + 1)].width = min(max(longest + 2, 8), 54)
        ws.freeze_panes = "A3"
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)   # a fresh instance has no deliverables/ yet
    wb.save(out); print("saved", out, "· sheets:", wb.sheetnames)

def write_csvs(sheets, outdir, stamp):
    os.makedirs(outdir or ".", exist_ok=True)   # a fresh instance has no deliverables/ yet
    for tab, header, rows, src in sheets:
        path = os.path.join(outdir, re.sub(r"[^\w.-]+", "-", tab.lower()) + ".csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(header)
            for r in rows:
                w.writerow(r)
            w.writerow([]); w.writerow(["# source: %s · rendered %s (to-table)" % (src, stamp)])
        print("saved", path, "(%d rows)" % len(rows))

def main(argv=None):
    ap = argparse.ArgumentParser(description="to-table generic renderer")
    ap.add_argument("instance")
    ap.add_argument("--section", action="append", default=[], help='"relpath[#anchor]:Tab Name"')
    ap.add_argument("--no-registers", action="store_true")
    ap.add_argument("--out", help="xlsx output path")
    ap.add_argument("--outdir", help="dir for csv outputs (default INSTANCE/deliverables)")
    ap.add_argument("--format", choices=["xlsx", "csv"], default="xlsx")
    ap.add_argument("--stamp", default="", help="render date (caller supplies; blank = unstamped)")
    a = ap.parse_args(argv)
    sheets = load_sources(a.instance, a.section, not a.no_registers)
    outdir = a.outdir or os.path.join(a.instance, "deliverables")
    if a.format == "xlsx":
        write_xlsx(sheets, a.out or os.path.join(outdir, "tables.xlsx"), a.stamp)
    else:
        write_csvs(sheets, outdir, a.stamp)

if __name__ == "__main__":
    sys.exit(main())
