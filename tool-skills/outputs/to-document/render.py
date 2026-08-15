#!/usr/bin/env python3
"""to-document — GENERIC renderer: a markdown content file -> a styled, neutral .docx. Instance-
agnostic: no product data lives here, only the house-neutral styling — so it travels with the
framework. The *content* markdown is authored per instance (by the agent, per ADAPTER.md); this
tool only applies the formatting. Requires python-docx (a docx library is why this is code).

Usage:  render.py CONTENT.md [--out FILE]

Supported markdown: '# Title', '> subtitle / meta' lines under the title, '## Section', paragraphs,
'- bullets' (with an optional '**Label** — ...' lead), pipe tables, '---' rules, and a trailing
'> footer: ...' line. Inline '**bold**' is honored.
"""
import argparse, os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x14, 0x17, 0x1D); ACCENT = RGBColor(0x1D, 0x46, 0xD6)
MUTED = RGBColor(0x5B, 0x64, 0x70); FONT = "Calibri"

def add_runs(p, text):
    """Render inline **bold** into a paragraph."""
    for k, seg in enumerate(text.split("**")):
        if seg:
            r = p.add_run(seg); r.bold = (k % 2 == 1)

def rule(doc, space_after=8):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); btm = OxmlElement("w:bottom")
    btm.set(qn("w:val"), "single"); btm.set(qn("w:sz"), "14"); btm.set(qn("w:space"), "1")
    btm.set(qn("w:color"), "1D46D6"); bdr.append(btm); pPr.append(bdr)

def heading(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = ACCENT; r.font.name = FONT
    sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "30"); r._element.get_or_add_rPr().append(sp)

def shade(cell, hexfill):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(sh)

def setcell(cell, text, bold=False, color=INK, white=False):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9.5); r.font.name = FONT
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if white else color

def emit_table(doc, block):
    header = [c.strip() for c in block[0].strip().strip("|").split("|")]
    body = [[c.strip().replace("**", "") for c in r.strip().strip("|").split("|")] for r in block[2:]]
    t = doc.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = True
    for c, h in zip(t.rows[0].cells, header):
        setcell(c, h, bold=True, white=True); shade(c, "1D46D6")
    for i, row in enumerate(body):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            setcell(cells[j], val, bold=(j == 0), color=ACCENT if j == 0 else INK)
        if i % 2 == 0:
            for cc in cells:
                shade(cc, "F3F5F9")

def render(md, out):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.9); s.bottom_margin = Cm(1.9); s.left_margin = Cm(2.1); s.right_margin = Cm(2.1)
    nm = doc.styles["Normal"]; nm.font.name = FONT; nm.font.size = Pt(10.5); nm.font.color.rgb = INK
    nm.paragraph_format.space_after = Pt(4); nm.paragraph_format.line_spacing = 1.18

    lines = md.splitlines(); i = 0; footer = None
    while i < len(lines):
        line = lines[i]; s = line.strip()
        if s.startswith("# "):                                   # title
            t = doc.add_paragraph(); t.paragraph_format.space_after = Pt(0)
            r = t.add_run(s[2:].strip()); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = INK
            i += 1; first = True
            while i < len(lines) and lines[i].strip().startswith(">"):   # subtitle / meta lines
                sub = lines[i].strip()[1:].strip()
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
                rr = p.add_run(sub); rr.font.color.rgb = MUTED
                rr.font.size = Pt(12) if first else Pt(8.5)
                if not first:
                    rr.font.name = "Consolas"
                first = False; i += 1
            rule(doc); continue
        if s.startswith("## "):
            heading(doc, s[3:].strip()); i += 1; continue
        if s == "---":
            rule(doc, space_after=4); i += 1; continue
        if s.lower().startswith("> footer:"):
            footer = s.split(":", 1)[1].strip(); i += 1; continue
        if s.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|:- "):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            emit_table(doc, block); continue
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
            add_runs(p, s[2:].strip()); i += 1; continue
        if s == "":
            i += 1; continue
        # paragraph: gather contiguous plain lines
        buf = [s]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|>|\||-\s|---)", lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        add_runs(doc.add_paragraph(), " ".join(buf))

    if footer:
        fp = doc.sections[0].footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rf = fp.add_run(footer); rf.font.size = Pt(7.5); rf.font.color.rgb = MUTED; rf.font.name = "Consolas"
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)   # a fresh instance has no deliverables/ yet
    doc.save(out); print("saved", out)

def main(argv=None):
    ap = argparse.ArgumentParser(description="to-document generic renderer (markdown -> styled .docx)")
    ap.add_argument("content"); ap.add_argument("--out")
    a = ap.parse_args(argv)
    out = a.out or os.path.splitext(a.content)[0] + ".docx"
    render(open(a.content, encoding="utf-8").read(), out)

if __name__ == "__main__":
    sys.exit(main())
